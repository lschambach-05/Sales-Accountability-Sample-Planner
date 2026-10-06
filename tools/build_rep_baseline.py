"""Build the rep baseline workbook (prior-year units/groups, retention) from the yearly MFP pull CSV.

Usage: python tools/build_rep_baseline.py <MFP pull .csv> <output .xlsx>
Rules are documented in docs/05-mfp-yearly-pull.md. Season names are currently 2024/2025;
change SEASONS_2025 and the 2024/2025 references for a new year.
"""
import sys, csv, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC, OUT = sys.argv[1], sys.argv[2]
R = list(csv.DictReader(open(SRC)))

CURRENT = ['BPR', 'JK', 'JMK', 'KJP', 'KJS', 'RB']
ROLLUP = {'BLS': 'KJS', 'LBS': 'KJP', 'LMD': 'JMK'}
PROGRAMS = {'Braided Pastry': 'Butter Braid Pastry', 'Combo': 'Combo', 'Wooden Spoon': 'Wooden Spoon CD',
            'Joyful Traditions': 'Joyful Tradition', 'Bella Napoli': 'Bella Napoli', 'Croissant Crown': 'Croissant Crowns',
            # Batavia Music Buffs runs all products under its own program (Lynwood, Oct 2026); MFP names it by year.
            'Batavia Music Buffs 2023': 'Batavia Music Buffs', 'Batavia Music Buffs 2025': 'Batavia Music Buffs'}
PROG_ORDER = ['Butter Braid Pastry', 'Combo', 'Wooden Spoon CD', 'Joyful Tradition', 'Bella Napoli', 'Croissant Crowns',
              'Batavia Music Buffs', 'Other (review)']
# Duplicate MFP group records for the same organization: {duplicate Group ID: Group ID to keep}.
GROUP_MERGES = {'199961': '84951'}   # St. Paul's Lutheran School (JK) - confirmed same school, Oct 2026
# Groups whose last owner is no longer a rep: {Group ID: current rep}. Filled in from the location match.
OWNER_OVERRIDES = {}

for r in R:
    r['Original Group ID'] = r['Group ID']
    r['Group ID'] = GROUP_MERGES.get(r['Group ID'], r['Group ID'])
SEASONS_2025 = ['Spring 2025', 'Fall 2025']

def rep_of(code):
    code = ROLLUP.get(code, code)
    return code if code in CURRENT else 'Unassigned'

def units(r): return float(r['Units'] or 0)

# ---- which fundraisers count ----
counted, excluded = [], []
for r in R:
    if r['Status'] == 'Canceled':
        excluded.append((r, 'Canceled')); continue
    if units(r) <= 0:
        excluded.append((r, 'Open/closed with no units sold' if r['Status'] == 'Open' else 'Closed with no units')); continue
    counted.append(r)

# ---- current owner of each group = owner of its most recent counted fundraiser ----
latest = {}
for r in counted:
    key = (r['Delivery Date'] or r['Start Date'], r['Fundraiser ID'])
    g = r['Group ID']
    if g not in latest or key > latest[g][0]:
        latest[g] = (key, r['Owning User'])
owner = {g: OWNER_OVERRIDES.get(g, rep_of(v[1])) for g, v in latest.items()}

# ---- years each group ran ----
ran_years = collections.defaultdict(set)
for r in counted:
    ran_years[r['Group ID']].add(int(r['Season'].split()[1]))

# ---- group x season x program rows ----
gsp = collections.OrderedDict()
for r in sorted(counted, key=lambda r: (r['Season'].split()[1], r['Season'].split()[0] == 'Fall', r['Group ID'])):
    prog = PROGRAMS.get(r['Program'], 'Other (review)')
    k = (r['Season'], r['Group ID'], prog)
    d = gsp.setdefault(k, {'units': 0.0, 'n': 0, 'codes': set(), 'name': r['Group Name'], 'mfp_prog': set(), 'flags': set()})
    d['units'] += units(r); d['n'] += 1; d['codes'].add(r['Owning User']); d['mfp_prog'].add(r['Program'])
    if r['Status'] == 'Open': d['flags'].add('Includes an Open (invoiced) fundraiser')
    if r['Original Group ID'] != r['Group ID']: d['flags'].add(f"Merged from duplicate Group ID {r['Original Group ID']}")
    if r['Group ID'] in OWNER_OVERRIDES: d['flags'].add('Rep assigned by location match')

seen_gs = set()
detail = []
for (season, gid, prog), d in gsp.items():
    yr = int(season.split()[1])
    prior_ran = (yr - 1) in ran_years[gid]
    retained = (1 if prior_ran else 0) if yr == 2025 else None
    first = 0 if (season, gid) in seen_gs else 1
    seen_gs.add((season, gid))
    detail.append([season, yr, season.split()[0], owner[gid], ', '.join(sorted(d['codes'])), int(gid), d['name'], prog,
                   ', '.join(sorted(d['mfp_prog'])), d['units'], d['n'], retained, first, '; '.join(sorted(d['flags']))])

# ---- styles ----
F = 'Arial'
hdr_font = Font(name=F, bold=True, color='FFFFFF'); hdr_fill = PatternFill('solid', fgColor='1F4E78')
body = Font(name=F); bold = Font(name=F, bold=True); title = Font(name=F, bold=True, size=14)
thin = Side(style='thin', color='BFBFBF')
def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c); cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()

# ================= Read Me =================
ws = wb.active; ws.title = 'Read Me'
ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 110
lines = [
    ('Rite Bite 2025 Rep Baseline (from MFP)', None),
    ('Source', 'MFP group sales 2024-2025.csv, pulled read-only from My Fundraising Place on 2026-10-05 (Sales > Fundraisers, Spring 2024 - Fall 2025). Totals match the pull\'s Season Check exactly.'),
    ('Seasons', 'By delivery date (from the pull): Spring = delivered Jan 1 - Jun 30, Fall = delivered Jul 1 - Dec 31.'),
    ('What counts', 'Closed fundraisers, plus Open fundraisers that were invoiced (units sold). Excluded: Canceled, and Open bookings with no units sold (all are well past their delivery date). Decision: Lynwood, 2026-10-06.'),
    ('Retained group', 'A 2025 group is retained if it ran ANY program in EITHER season of 2024. Its units count toward the program it ran in 2025. Decision: Lynwood. 2024 retention is not shown because 2023 was not pulled.'),
    ('Current rep', 'Each group\'s whole history is credited to its current rep = the owning user on its most recent counted fundraiser, after rollups. Decision: Lynwood.'),
    ('Rep rollups', 'BLS -> KJS, LBS -> KJP, LMD -> JMK (from the MFP pull). Current reps: BPR, JK, JMK, KJP, KJS, RB (Lynwood). BJS, GLP and LSS were not assigned, so groups last owned by them show as "Unassigned".'),
    ('Programs', 'MFP Braided Pastry = Butter Braid Pastry; Combo is its own program; Batavia Music Buffs is its own program and runs all products (Lynwood). Wooden Spoon -> Wooden Spoon CD, Joyful Traditions -> Joyful Tradition, Croissant Crown -> Croissant Crowns. Anything else -> "Other (review)".'),
    ('Prior Year', 'For a 2025 row: same rep, same program, same season in 2024.'),
    ('Read with care', 'Program-level retention can exceed 100%: groups that switch programs still count as retained. The Rep Totals tab is the cleanest retention view.'),
    ('Group matching', 'Groups are matched on MFP Group ID. Confirmed duplicates are merged: 199961 -> 84951 (St. Paul\'s Lutheran School, JK). Other same-name IDs are listed on Review and NOT merged.'),
    ('Tabs', 'Rep Program Summary: planner-ready rows (rep x program x season). Rep Totals: per rep per season. Group Detail: one row per group x program x season (the data every formula reads). Review: items to check.'),
]
for i, (a, b) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=a).font = title if b is None else bold
    if b: c = ws.cell(row=i, column=2, value=b); c.font = body; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=i, column=1).alignment = Alignment(vertical='top')

# ================= Group Detail =================
gd = wb.create_sheet('Group Detail')
gcols = ['Season', 'Year', 'Season Name', 'Current Rep', 'MFP Owning User(s)', 'Group ID', 'Group Name', 'Program',
         'Program in MFP', 'Units', 'Fundraisers', 'Retained (ran 2024)', 'Count Group (1st row this season)', 'Notes']
header(gd, 1, gcols, [12, 7, 9, 10, 14, 10, 44, 20, 22, 9, 11, 12, 13, 36])
for i, row in enumerate(detail, 2):
    for j, v in enumerate(row, 1):
        c = gd.cell(row=i, column=j, value=v); c.font = body
    gd.cell(row=i, column=10).number_format = '#,##0'
N = len(detail) + 1
gd.freeze_panes = 'A2'
gd.auto_filter.ref = f'A1:N{N}'
GD = lambda col: f"'Group Detail'!${col}$2:${col}${N}"
S_, REP, PROG, UN, RET, FIRST = GD('A'), GD('D'), GD('H'), GD('J'), GD('L'), GD('M')

# ================= Rep Program Summary =================
ps = wb.create_sheet('Rep Program Summary', 1)
pcols = ['Season', 'Prior Year Season', 'Rep', 'Program', 'Prior Year Units', 'Prior Year Groups', 'Units', 'Groups',
         'Retained Units', 'Retained Groups', 'New Units', 'New Groups', 'Unit Retention %', 'Group Retention %',
         'Avg Units per New Group', 'Avg Units per Retained Group']
header(ps, 1, pcols, [12, 11, 11, 20, 11, 10, 10, 9, 10, 10, 10, 9, 10, 10, 11, 11])
present = {(d[0], d[3], d[7]) for d in detail}
row = 2
for season in SEASONS_2025:
    prior = season.replace('2025', '2024')
    for rep in CURRENT + ['Unassigned']:
        for prog in PROG_ORDER:
            if (season, rep, prog) not in present and (prior, rep, prog) not in present: continue
            r = row
            vals = [season, prior, rep, prog,
                    f'=SUMIFS({UN},{S_},$B{r},{REP},$C{r},{PROG},$D{r})',
                    f'=COUNTIFS({S_},$B{r},{REP},$C{r},{PROG},$D{r})',
                    f'=SUMIFS({UN},{S_},$A{r},{REP},$C{r},{PROG},$D{r})',
                    f'=COUNTIFS({S_},$A{r},{REP},$C{r},{PROG},$D{r})',
                    f'=SUMIFS({UN},{S_},$A{r},{REP},$C{r},{PROG},$D{r},{RET},1)',
                    f'=COUNTIFS({S_},$A{r},{REP},$C{r},{PROG},$D{r},{RET},1)',
                    f'=G{r}-I{r}', f'=H{r}-J{r}',
                    f'=IF(E{r}=0,"",I{r}/E{r})', f'=IF(F{r}=0,"",J{r}/F{r})',
                    f'=IF(L{r}=0,"",K{r}/L{r})', f'=IF(J{r}=0,"",I{r}/J{r})']
            for j, v in enumerate(vals, 1):
                c = ps.cell(row=r, column=j, value=v); c.font = body
            for col in 'EFGHIJKL': ps[f'{col}{r}'].number_format = '#,##0;(#,##0);-'
            for col in 'MN': ps[f'{col}{r}'].number_format = '0.0%'
            for col in 'OP': ps[f'{col}{r}'].number_format = '#,##0.0'
            if rep == 'Unassigned' or prog == 'Other (review)':
                for j in range(1, 17): ps.cell(row=r, column=j).fill = PatternFill('solid', fgColor='FFF2CC')
            row += 1
last = row - 1
ps.cell(row=row, column=1, value='Total').font = bold
for col in 'EFGHIJKL':
    c = ps[f'{col}{row}']; c.value = f'=SUM({col}2:{col}{last})'; c.font = bold; c.number_format = '#,##0;(#,##0);-'
ps.cell(row=row + 2, column=1, value='Yellow rows need review (Unassigned rep or unrecognised program). Group counts in the Total row add up program rows, so a group running two programs counts twice; use Rep Totals for distinct groups.').font = Font(name=F, italic=True)
ps.freeze_panes = 'E2'; ps.auto_filter.ref = f'A1:P{last}'

# ================= Rep Totals =================
rt = wb.create_sheet('Rep Totals', 2)
tcols = ['Season', 'Prior Year Season', 'Rep', 'Prior Year Units', 'Prior Year Groups', 'Units', 'Groups',
         'Retained Units', 'Retained Groups', 'New Units', 'New Groups', 'Unit Retention %', 'Group Retention %',
         'Avg Units per New Group', 'Avg Units per Retained Group']
header(rt, 1, tcols, [12, 11, 11, 11, 10, 10, 9, 10, 10, 10, 9, 10, 10, 11, 11])
row = 2
for season in SEASONS_2025:
    prior = season.replace('2025', '2024'); start = row
    for rep in CURRENT + ['Unassigned']:
        r = row
        vals = [season, prior, rep,
                f'=SUMIFS({UN},{S_},$B{r},{REP},$C{r})',
                f'=SUMIFS({FIRST},{S_},$B{r},{REP},$C{r})',
                f'=SUMIFS({UN},{S_},$A{r},{REP},$C{r})',
                f'=SUMIFS({FIRST},{S_},$A{r},{REP},$C{r})',
                f'=SUMIFS({UN},{S_},$A{r},{REP},$C{r},{RET},1)',
                f'=SUMIFS({FIRST},{S_},$A{r},{REP},$C{r},{RET},1)',
                f'=F{r}-H{r}', f'=G{r}-I{r}',
                f'=IF(D{r}=0,"",H{r}/D{r})', f'=IF(E{r}=0,"",I{r}/E{r})',
                f'=IF(K{r}=0,"",J{r}/K{r})', f'=IF(I{r}=0,"",H{r}/I{r})']
        for j, v in enumerate(vals, 1): rt.cell(row=r, column=j, value=v).font = body
        if rep == 'Unassigned':
            for j in range(1, 16): rt.cell(row=r, column=j).fill = PatternFill('solid', fgColor='FFF2CC')
        row += 1
    r = row
    rt.cell(row=r, column=1, value=season).font = bold
    rt.cell(row=r, column=3, value='All reps').font = bold
    for col in 'DEFGHIJK': rt[f'{col}{r}'] = f'=SUM({col}{start}:{col}{r-1})'
    for col, f in zip('LMNO', [f'=IF(D{r}=0,"",H{r}/D{r})', f'=IF(E{r}=0,"",I{r}/E{r})', f'=IF(K{r}=0,"",J{r}/K{r})', f'=IF(I{r}=0,"",H{r}/I{r})']):
        rt[f'{col}{r}'] = f
    for j in range(1, 16): rt.cell(row=r, column=j).font = bold; rt.cell(row=r, column=j).border = Border(top=thin)
    row += 2
for rr in rt.iter_rows(min_row=2, max_row=row):
    for c in rr:
        if c.column in range(4, 12): c.number_format = '#,##0;(#,##0);-'
        elif c.column in (12, 13): c.number_format = '0.0%'
        elif c.column in (14, 15): c.number_format = '#,##0.0'
rt.freeze_panes = 'D2'

# ================= Review =================
rv = wb.create_sheet('Review')
header(rv, 1, ['Check', 'Group ID', 'Group Name', 'Detail', 'Units affected'], [34, 10, 44, 80, 12])
items = []
for gid, rep in sorted(owner.items(), key=lambda x: x[0]):
    if rep == 'Unassigned':
        g = [d for d in detail if d[5] == int(gid)]
        items.append(('Unassigned rep (last owner not a current rep)', int(gid), g[0][6], 'Last owner: ' + latest[gid][1] + '. Seasons: ' + ', '.join(sorted({d[0] for d in g})), sum(d[9] for d in g)))
for d in detail:
    if d[7] == 'Other (review)':
        items.append(('Unrecognised program', d[5], d[6], f'{d[0]}: MFP program "{d[8]}", rep {d[3]}', d[9]))
for r, why in excluded:
    if r['Status'] == 'Open':
        items.append(('Excluded Open booking (nothing sold)', int(r['Group ID']), r['Group Name'], f"{r['Season']}, {r['Owning User']}, state {r['Fundraiser State']}, delivery {r['Delivery Date'] or 'none'}", 0))
moved = collections.defaultdict(set)
for r in counted: moved[r['Group ID']].add(rep_of(r['Owning User']))
for gid, reps in moved.items():
    reps.discard('Unassigned')
    if len(reps) > 1:
        g = [d for d in detail if d[5] == int(gid)]
        items.append(('Group moved between current reps', int(gid), g[0][6], 'Credited to ' + owner[gid] + '; also ran under ' + ', '.join(sorted(reps - {owner[gid]})), sum(d[9] for d in g)))
# same-name, same-rep duplicate IDs that would flip a 2025 group from new to retained
byname = collections.defaultdict(set)
names = {}
for d in detail: byname[(d[6].strip().lower(), d[3])].add(d[5]); names[d[5]] = d[6]
for (nm, rep), ids in byname.items():
    if len(ids) < 2: continue
    for gid in ids:
        new25 = [d for d in detail if d[5] == gid and d[1] == 2025 and d[11] == 0]
        sib_ran = any(2024 in ran_years[str(o)] for o in ids if o != gid)
        detail_txt = f'Same name, same rep ({rep}) as Group ID(s) ' + ', '.join(str(o) for o in sorted(ids - {gid}))
        if new25 and sib_ran:
            items.append(('Possible duplicate group record: would be RETAINED if merged', gid, names[gid], detail_txt + '; the other ID ran in 2024', sum(d[9] for d in new25)))
        else:
            items.append(('Same-name group IDs (no effect on 2025 retention)', gid, names[gid], detail_txt, 0))
items.append(('Excluded: Canceled fundraisers', None, None, f"{sum(1 for _, w in excluded if w == 'Canceled')} canceled fundraisers left out", 0))
for i, it in enumerate(items, 2):
    for j, v in enumerate(it, 1):
        c = rv.cell(row=i, column=j, value=v); c.font = body
    rv.cell(row=i, column=5).number_format = '#,##0;(#,##0);-'
rv.freeze_panes = 'A2'; rv.auto_filter.ref = f'A1:E{len(items)+1}'

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('detail rows', len(detail), 'summary rows', last - 1, 'review items', len(items))
print('counted units by season', {s: sum(units(r) for r in counted if r['Season'] == s) for s in ['Spring 2024','Fall 2024','Spring 2025','Fall 2025']})
print('excluded', collections.Counter(w for _, w in excluded))
